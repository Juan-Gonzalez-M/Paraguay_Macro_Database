#!/usr/bin/env python3
"""Versión 2 del calendario de eventos del usuario: agrega los eventos que faltaban, en las mismas hojas y en el mismo informe.

Uso (desde esta carpeta; requiere openpyxl y python-docx, p. ej. el entorno ~/.venvs/edh):
    ~/.venvs/edh/bin/python agregar_eventos.py

Entradas (no se modifican; SHA-256 en inventory.csv):
    raw/BCP_eventos_politica_2011_2026.xlsx   libro del usuario, tal como lo entregó el 2026-09-25
    raw/Informe_eventos_BCP_2011_2026.docx    informe del usuario, ídem
Salidas:
    BCP_eventos_politica_2011_2026_v2.xlsx    mismas hojas y columnas; filas nuevas al final de cada hoja
    Informe_eventos_BCP_2011_2026_v2.docx     mismo texto; se agrega el «Anexo C» con los eventos nuevos

Reglas:
  - Ninguna fila existente se modifica ni se borra. Cada fila nueva lleva su tipo de evidencia en verification_status:
      verificado_fuente_BCP_archivo  comunicado del BCP leído en una copia `id_` del Internet Archive (SHA-256 en
                                     ../web_brechas_2026-09-23/inventory_archivo.csv);
      inferido_serie_diaria_BCP      cambio deducido de las tasas diarias FPD/FPL de la base (fuente: BCP, mercado
                                     interbancario) y del calendario del CPM; SIN documento que lo respalde.
  - Los valores que el comunicado no publica (porcentaje de encaje, penalidades de LRM) quedan vacíos (la Guía del libro:
    vacío = no extraído, nunca cero).
  - La familia «Corredor» es nueva: facilidades permanentes (FPD/FPL) y LRM, separada de las decisiones de TPM.
"""
import copy, os
import openpyxl, docx

AQUI = os.path.dirname(os.path.abspath(__file__))
WB_IN, WB_OUT = "raw/BCP_eventos_politica_2011_2026.xlsx", "BCP_eventos_politica_2011_2026_v2.xlsx"
DOC_IN, DOC_OUT = "raw/Informe_eventos_BCP_2011_2026.docx", "Informe_eventos_BCP_2011_2026_v2.docx"

WB16 = "https://web.archive.org/web/20220711042210id_/https://www.bcp.gov.py/userfiles/files/01Comunicado_primeras_medidas_16_03_20%281%29.pdf"
WB30 = "https://web.archive.org/web/20230204000724id_/https://www.bcp.gov.py/userfiles/files/05Comunicado_Medidas_adicionales_30_03_20(1).pdf"
SERIE = "base DuckDB (main.v_series_research): ib_facilidad_permanente_de_deposito_fpd_tasa_pct e ib_facilidad_permanente_de_liquidez_fpl_tasa_pct; cambios en proyectos/25_n4_sorpresas_monetarias/datos/cambios_corredor_inferidos.csv; TPM del calendario del CPM (web_brechas_2026-09-23)"
V_ARCH, V_INF = "verificado_fuente_BCP_archivo", "inferido_serie_diaria_BCP"
INTERP_INF = "Inferido de las tasas diarias; falta la resolución o el comunicado que lo respalde. No usar como evento documentado sin verificar."

# Eventos: event_id, announcement_date, approval_date, effective_date, end_date, family, subfamily, event_type, shock_type,
#          sector, norm, title, description, stated_rationale, researcher_interpretation, verification_status
EVENTOS = [
    ["REL_2020_0316_PKG", "2020-03-16", "", "", "", "Alivio", "Prudencial/liquidez", "anuncio_paquete", "COVID-19", "sistema financiero", "",
     "Primeras medidas COVID (comunicado del 16-03)",
     "Comunicado del BCP del 16-03-2020: normativas que amplían el régimen para formalizar renovaciones, refinanciaciones y reestructuraciones; extensión del plazo de enajenación de bienes adjudicados en pago (posterga previsiones); uso de un porcentaje del encaje legal para liquidez; rebaja de la FPL y de sus tramos; menores penalidades por cancelación anticipada de LRM.",
     "Aliviar los efectos del COVID-19 sobre familias y empresas (texto del comunicado).",
     "El comunicado del 16-03 dice que las resoluciones ya fueron emitidas; REL_2020_REG fecha la Res. 4/Acta 18 el 18-03. Verificar si son los mismos actos antes de fechar el shock.",
     V_ARCH],
    ["RES_2020_0316_USE", "2020-03-16", "", "", "", "Encaje", "Requisitos", "alivio", "", "", "",
     "Uso de encaje para liquidez de clientes",
     "«Se ha determinado un porcentaje [de los fondos de encaje legal] para aplicar dicha liquidez a satisfacer requerimientos de clientes del sistema financiero». El comunicado no publica el porcentaje ni la resolución.",
     "", "Porcentaje, moneda y vigencia pendientes de la resolución.", V_ARCH],
    ["COR_2020_0316_FPL", "2020-03-16", "", "2020-03-17", "", "Corredor", "Facilidades permanentes", "parámetro", "COVID-19", "", "",
     "Rebaja de la FPL y de sus tramos",
     "FPL a 1 día (reporto): −100 pb, de 4,50% a 3,50%. FPL Primer Tramo (hasta 30 días): de TPM + 200 pb a TPM + 75 pb. FPL Segundo Tramo (30 días adicionales): de TPM + 300 pb a TPM + 125 pb.",
     "Flexibilizar las condiciones de las ventanillas de liquidez (texto del comunicado).",
     "Serie diaria de la FPL: 4,50 (13-03) → 4,25 (16-03, por la TPM del 13-03) → 3,50 (17-03). El «de 4,50% a 3,50%» combina ese paso con la TPM de la extraordinaria del 16-03 (3,25%): desde el 17-03 la FPL queda en TPM + 25 pb (antes TPM + 50 pb).",
     V_ARCH],
    ["COR_2020_0316_LRM", "2020-03-16", "", "", "", "Corredor", "LRM", "parámetro", "COVID-19", "", "",
     "Menor penalidad por cancelación anticipada de LRM",
     "El BCP disminuyó las tasas de penalización por cancelación anticipada de Letras de Regulación Monetaria. El comunicado no publica los valores.",
     "Permitir que las entidades hagan líquidos los títulos sin mayores costos (texto del comunicado).", "Valores pendientes de la resolución.", V_ARCH],
    ["RES_2020_0330_RELEASE", "2020-03-30", "", "", "", "Encaje", "Requisitos", "alivio", "", "", "",
     "Liberación de encaje MN y ME",
     "Comunicado del BCP del 30-03-2020: «Liberación del Encaje Legal en Moneda Nacional y Extranjera», fondos por un valor equivalente a USD 740 millones.",
     "Asegurar la liquidez y la concesión de créditos (texto del comunicado).",
     "Mismo día que la segunda reunión extraordinaria del CPM (TPM −100 pb, a 2,25%; ver el calendario del CPM). Tasas y tramos pendientes de la resolución.",
     V_ARCH],
    ["REL_2020_0330_WINDOW", "2020-03-30", "", "", "", "Alivio", "Prudencial/liquidez", "instrumento_adicional", "COVID-19", "intermediarios", "",
     "Ventanilla de liquidez hasta 12 meses",
     "Comunicado del BCP del 30-03-2020: «Creación de una ventanilla de liquidez para las entidades por un plazo de hasta 12 meses, por un valor equivalente a USD 760 millones».",
     "Asegurar la liquidez y la concesión de créditos (texto del comunicado).",
     "Probablemente es la misma facilidad que REL_2020_FCE (Res. 1/Acta 21, misma fecha): relación «mismo_acto_probable». No sumar ambos como dos shocks.",
     V_ARCH],
    ["COR_2013_0411_FPD", "", "", "2013-04-11", "", "Corredor", "Facilidades permanentes", "parámetro_inferido", "", "", "",
     "FPD fijada en TPM − 100 pb",
     "La FPD pasa de 5,47% a 4,50% con la TPM en 5,50%. Antes fluctuaba entre 5,21% y 5,50% (cerca de la TPM); desde entonces se mueve con la TPM a −100 pb.",
     "", INTERP_INF, V_INF],
    ["COR_2015_0320", "", "", "2015-03-20", "", "Corredor", "Facilidades permanentes", "parámetro_inferido", "", "", "",
     "Corredor de ±100 a ±75 pb",
     "El 19-03 la TPM baja de 6,75% a 6,50%; al día hábil siguiente la FPL baja 50 pb (7,75% → 7,25%) y la FPD no se mueve (5,75%). El corredor pasa de TPM ± 100 pb a TPM ± 75 pb.",
     "", INTERP_INF, V_INF],
    ["COR_2015_0916", "", "", "2015-09-16", "", "Corredor", "Facilidades permanentes", "parámetro_inferido", "", "", "",
     "Corredor asimétrico: TPM − 25 / + 100 pb",
     "Sin cambio de TPM (5,75% desde julio): la FPD sube de 5,00% a 5,50% y la FPL de 6,50% a 6,75%. El corredor pasa de TPM ± 75 pb a TPM − 25 / + 100 pb. Ocurre el día anterior a la reunión del CPM del 17-09-2015.",
     "", INTERP_INF, V_INF],
    ["COR_2019_0610", "", "", "2019-06-10", "", "Corredor", "Facilidades permanentes", "parámetro_inferido", "", "", "",
     "FPL de TPM + 100 a TPM + 50 pb",
     "Sin cambio de TPM (4,75%): la FPL baja de 5,75% a 5,25%; la FPD no se mueve. Un día antes del paquete del 11-06-2019 (REL_2019_RESERVE, REL_2019_ASSETS).",
     "", INTERP_INF, V_INF],
    ["COR_2020_0323", "", "", "2020-03-23", "2020-03-24", "Corredor", "Facilidades permanentes", "parámetro_inferido", "COVID-19", "", "",
     "Salto de un día de la FPD",
     "La FPD sube de 3,00% a 3,50% el 23-03-2020 y vuelve a 3,00% el 24-03; la FPL no se mueve. No hay decisión conocida que lo explique.",
     "", INTERP_INF + " Puede ser un error de la serie.", V_INF],
]
RELACIONES = [  # event_id, related_event_id, relation_type, note
    ["REL_2020_0316_PKG", "REL_2020_REG", "posible_mismo_acto", "Comunicado del 16-03 vs. Res. 4/Acta 18 del 18-03: verificar."],
    ["RES_2020_0316_USE", "REL_2020_0316_PKG", "mismo_paquete", ""],
    ["COR_2020_0316_FPL", "REL_2020_0316_PKG", "mismo_paquete", ""],
    ["COR_2020_0316_LRM", "REL_2020_0316_PKG", "mismo_paquete", ""],
    ["RES_2020_0330_RELEASE", "REL_2020_FCE", "mismo_paquete", "Comunicado del 30-03-2020 (medidas adicionales)."],
    ["REL_2020_0330_WINDOW", "REL_2020_FCE", "mismo_acto_probable", "Ventanilla hasta 12 meses por USD 760 millones; no duplicar el shock."],
]
PARAMETROS = [  # event_id, parameter, old_value, new_value, unit, scope, note
    ["COR_2020_0316_FPL", "tasa_fpl_1_dia", "4.5", "3.5", "porcentaje", "FPL a 1 día", "Según el comunicado."],
    ["COR_2020_0316_FPL", "spread_fpl_tramo1", "200", "75", "puntos básicos sobre la TPM", "FPL Primer Tramo (hasta 30 días)", ""],
    ["COR_2020_0316_FPL", "spread_fpl_tramo2", "300", "125", "puntos básicos sobre la TPM", "FPL Segundo Tramo", ""],
    ["COR_2020_0316_FPL", "spread_fpl", "50", "25", "puntos básicos sobre la TPM", "FPL a 1 día", "Inferido de la serie diaria (17-03-2020)."],
    ["RES_2020_0330_RELEASE", "monto_liberado", "", "740", "USD millones (equivalente)", "MN y ME", "Según el comunicado."],
    ["REL_2020_0330_WINDOW", "monto_ventanilla", "", "760", "USD millones (equivalente)", "hasta 12 meses", "Según el comunicado."],
    ["COR_2013_0411_FPD", "spread_fpd", "", "-100", "puntos básicos sobre la TPM", "FPD", "Inferido; antes fluctuaba cerca de la TPM."],
    ["COR_2015_0320", "spread_fpd", "-100", "-75", "puntos básicos sobre la TPM", "FPD", "Inferido."],
    ["COR_2015_0320", "spread_fpl", "100", "75", "puntos básicos sobre la TPM", "FPL", "Inferido."],
    ["COR_2015_0916", "spread_fpd", "-75", "-25", "puntos básicos sobre la TPM", "FPD", "Inferido."],
    ["COR_2015_0916", "spread_fpl", "75", "100", "puntos básicos sobre la TPM", "FPL", "Inferido."],
    ["COR_2019_0610", "spread_fpl", "100", "50", "puntos básicos sobre la TPM", "FPL", "Inferido."],
]
ALIVIOS = [  # event_id, shock_episode, eligible_sector, renewal, refinancing, restructuring, moratorium_treatment,
             # provision_relief, liquidity_facility, reserve_release, eligibility_end, verification_status, note
    ["REL_2020_0316_PKG", "COVID-19", "sectores afectados", "sí", "sí", "sí", "", "sí", "sí", "sí", "", V_ARCH,
     "Según el comunicado del 16-03-2020; previsiones: solo por bienes adjudicados en pago. Vacío = no extraído."],
    ["REL_2020_0330_WINDOW", "COVID-19", "intermediarios", "", "", "", "", "", "sí", "", "", V_ARCH, "Ventanilla hasta 12 meses."],
]
FUENTES = [  # event_id, source_type, source_url, evidence_scope, verification_status
    *[[e, "comunicado BCP (copia del Internet Archive)", WB16, "Texto completo del comunicado; SHA-256 en web_brechas_2026-09-23/inventory_archivo.csv", V_ARCH]
      for e in ("REL_2020_0316_PKG", "RES_2020_0316_USE", "COR_2020_0316_FPL", "COR_2020_0316_LRM")],
    *[[e, "comunicado BCP (copia del Internet Archive)", WB30, "Texto completo del comunicado; SHA-256 en web_brechas_2026-09-23/inventory_archivo.csv", V_ARCH]
      for e in ("RES_2020_0330_RELEASE", "REL_2020_0330_WINDOW")],
    *[[e, "serie diaria del BCP en la base (inferencia)", SERIE, "Cambio de tasas y spreads; sin documento", V_INF]
      for e in ("COR_2013_0411_FPD", "COR_2015_0320", "COR_2015_0916", "COR_2019_0610", "COR_2020_0323", "COR_2020_0316_FPL")],
]
GUIA = [
    ["verificado_fuente_BCP_archivo", "Agregado el 2026-09-25: comunicado del BCP leído en una copia del Internet Archive; el hash del archivo está en el inventario del repositorio."],
    ["inferido_serie_diaria_BCP", "Agregado el 2026-09-25: deducido de las tasas diarias FPD/FPL y del calendario del CPM; no hay documento. No es una verificación."],
    ["Familia Corredor", "Agregada el 2026-09-25: facilidades permanentes (FPD/FPL, tramos) y LRM, separadas de las decisiones de TPM (que están en el calendario del CPM)."],
]
POR_HOJA = {"Eventos": EVENTOS, "Relaciones": RELACIONES, "Parámetros": PARAMETROS, "Alivios": ALIVIOS, "Fuentes": FUENTES, "Guía": GUIA}


def ultima_fila(ws):
    n = ws.max_row
    while n > 1 and all(c.value in (None, "") for c in ws[n]):
        n -= 1
    return n


def agregar(ws, filas):
    ancho = max(c.column for c in ws[1] if c.value not in (None, ""))
    for f in filas:
        assert len(f) == ancho, (ws.title, f[0], len(f), ancho)
    n = ultima_fila(ws)
    ids = {ws.cell(r, 1).value for r in range(2, n + 1)}
    for i, f in enumerate(filas, start=1):
        if ws.title == "Eventos":
            assert f[0] not in ids, f"event_id repetido: {f[0]}"
        for j, v in enumerate(f, start=1):
            c = ws.cell(n + i, j, v if v != "" else None)
            m = ws.cell(n, j)
            if m.has_style:
                c._style = copy.copy(m._style)


def main():
    os.chdir(AQUI)
    wb = openpyxl.load_workbook(WB_IN)
    antes = {ws.title: ultima_fila(ws) for ws in wb.worksheets}
    ids = {wb["Eventos"].cell(r, 1).value for r in range(2, antes["Eventos"] + 1)}
    for hoja, filas in POR_HOJA.items():
        for f in filas:
            if hoja not in ("Eventos", "Guía"):
                assert f[0] in ids or f[0] in {e[0] for e in EVENTOS}, (hoja, f[0])
            if hoja == "Relaciones":
                assert f[1] in ids or f[1] in {e[0] for e in EVENTOS}, (hoja, f[1])
        agregar(wb[hoja], filas)
    wb.save(WB_OUT)
    wb2 = openpyxl.load_workbook(WB_OUT)
    for ws in wb2.worksheets:
        print(f"{ws.title}: {antes[ws.title]} -> {ultima_fila(ws)} filas (+{len(POR_HOJA.get(ws.title, []))})")
        assert ultima_fila(ws) == antes[ws.title] + len(POR_HOJA.get(ws.title, []))

    d = docx.Document(DOC_IN)
    d.add_heading("Anexo C. Eventos agregados en la verificación del 25-09-2026", level=1)
    d.add_paragraph(
        "Este anexo se agregó al verificar el libro con documentos y datos del repositorio. El texto anterior no se modificó. "
        "Las filas nuevas llevan su tipo de evidencia: «verificado_fuente_BCP_archivo» (comunicado del BCP en una copia del "
        "Internet Archive, con hash en el inventario) o «inferido_serie_diaria_BCP» (deducido de las tasas diarias FPD/FPL; "
        "sin documento, no es una verificación). Se crea la familia «Corredor» para las facilidades permanentes y las LRM.")
    d.add_paragraph(
        "Paquete COVID de marzo de 2020. El libro no registraba el comunicado del 16-03 ni las medidas de encaje de 2020. El 16-03 el BCP "
        "anunció normativas de renovación, refinanciación y reestructuración, más plazo para enajenar bienes adjudicados, el uso de un "
        "porcentaje del encaje para liquidez, la rebaja de la FPL a 1 día (de 4,50% a 3,50%) y de sus tramos (TPM + 75 y + 125 pb), y menores "
        "penalidades por cancelar LRM. El 30-03 anunció la liberación de encaje en MN y ME (unos USD 740 millones) y una ventanilla de "
        "liquidez de hasta 12 meses (USD 760 millones), probablemente la FCE ya registrada. El comunicado del 16-03 da por emitidas las "
        "resoluciones, mientras que el libro fecha la Res. 4/Acta 18 el 18-03: queda para verificar.")
    d.add_paragraph(
        "Estructura del corredor (inferida). Las tasas diarias muestran cambios de spread sin decisión de TPM o que no se explican solo por ella: "
        "FPD en TPM − 100 pb desde el 11-04-2013; corredor de ±100 a ±75 pb el 20-03-2015; de ±75 a −25/+100 pb el 16-09-2015; FPL de TPM + 100 "
        "a + 50 pb el 10-06-2019; FPL a TPM + 25 pb el 17-03-2020 (documentado); y un salto de un día de la FPD el 23-03-2020. Salvo el de 2020, "
        "falta el documento que respalde cada uno.")
    t = d.add_table(rows=1, cols=4)
    for c, h in zip(t.rows[0].cells, ["event_id", "Fecha", "Título", "Evidencia"]):
        c.text = h
    for e in EVENTOS:
        r = t.add_row().cells
        r[0].text, r[1].text, r[2].text, r[3].text = e[0], e[1] or e[3], e[11], e[15]
    d.save(DOC_OUT)
    print("docx:", DOC_OUT, "| eventos en el anexo:", len(EVENTOS))


if __name__ == "__main__":
    main()
